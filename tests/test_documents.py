import base64
import io

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader

from kernelledger.desktop.helper import Helper
from kernelledger.documents import export_document
from kernelledger.ledger import Ledger
from kernelledger.notebook import Notebook, source_hash
from kernelledger.source import SourceError


def document_text(path, format):
    if format == "pdf":
        return "\n".join(page.extract_text() for page in PdfReader(path).pages)
    document = Document(path)
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_offline_export_contains_current_source_recorded_outputs_and_stale_warning(
    tmp_path, format
):
    notebook = Notebook(tmp_path / "report")
    notebook.init()
    notebook.add(
        "markdown",
        "# Results\nRésumé αβ Привет\n\n- Baseline\n\n| Name | Value |\n| --- | --- |\n| Mean | 16.5 |",
        "intro",
    )
    notebook.add("code", "raise AssertionError('EXPORT MUST NEVER EXECUTE THIS')", "code")
    notebook.add("raw", "<untrusted> raw source & characters", "raw")
    ledger = Ledger(notebook.root / ".kl/ledger.sqlite")
    sid = ledger.session(notebook.read()["id"], {})
    eid = ledger.execution(sid, "code", source_hash("older source"))
    ledger.write("UPDATE executions SET status='error' WHERE exec_id=?", (eid,))
    ledger.event(eid, "stdout", {"text": "Recorded value: 42\n"})
    ledger.event(
        eid,
        "error",
        {"ename": "ValueError", "evalue": "recorded failure", "traceback": ["Traceback line"]},
    )
    image = io.BytesIO()
    Image.new("RGB", (24, 16), "blue").save(image, "PNG")
    ledger.event(
        eid,
        "display",
        {"bundle": {"image/png": {"value": base64.b64encode(image.getvalue()).decode()}}},
    )
    ledger.close()
    path = tmp_path / f"report.{format}"
    result = Helper(notebook.root).dispatch(
        {"method": "export_document", "format": format, "path": str(path)}
    )
    assert result == {"path": str(path), "format": format}
    text = document_text(path, format)
    for expected in (
        "Results",
        "Résumé",
        "αβ",
        "Привет",
        "Baseline",
        "Mean",
        "16.5",
        "EXPORT MUST NEVER EXECUTE THIS",
        "Recorded value: 42",
        "earlier version",
        "ValueError: recorded failure",
        "Traceback line",
        "<untrusted> raw source & characters",
    ):
        assert expected in text
    if format == "docx":
        assert len(Document(path).inline_shapes) == 1
        assert len(Document(path).tables) == 1
    else:
        assert sum(len(page.images) for page in PdfReader(path).pages) == 1
    assert not (notebook.root / ".kl/supervisor.json").exists()


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_export_without_history_and_long_cells(tmp_path, format):
    notebook = Notebook(tmp_path / "source-only")
    notebook.init()
    notebook.add("code", "\n".join(f"print('source line {i}')" for i in range(220)), "long")
    path = tmp_path / f"long.{format}"
    Helper(notebook.root).dispatch(
        {"method": "export_document", "format": format, "path": str(path)}
    )
    text = document_text(path, format)
    assert "source line 0" in text and "source line 219" in text
    if format == "pdf":
        assert len(PdfReader(path).pages) >= 3
    assert not (notebook.root / ".kl/ledger.sqlite").exists()


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_empty_notebook_export(tmp_path, format):
    notebook = Notebook(tmp_path / "empty")
    notebook.init()
    path = tmp_path / f"empty.{format}"
    export_document(notebook, str(path), format)
    assert "This notebook has no cells." in document_text(path, format)


def test_invalid_export_does_not_replace_destination(tmp_path):
    notebook = Notebook(tmp_path / "notebook")
    notebook.init()
    path = tmp_path / "existing.pdf"
    path.write_bytes(b"preserve me")
    with pytest.raises(SourceError, match="format"):
        export_document(notebook, str(path), "html")
    with pytest.raises(SourceError, match=".docx"):
        export_document(notebook, str(path), "docx")
    assert path.read_bytes() == b"preserve me"
