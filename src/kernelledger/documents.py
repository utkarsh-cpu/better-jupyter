"""Offline, paginated PDF and editable Word exports of source and recorded output."""

import base64
import io
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, cast
from xml.sax.saxutils import escape

from .ledger import fold
from .notebook import Notebook, source_hash
from .sharing import payload
from .source import SourceError, atomic


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    data: bytes = b""
    rows: list[list[str]] | None = None


def clean(text: str) -> str:
    # XML 1.0 controls and terminal ANSI sequences are not document content.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", re.sub(r"\x1b\[[0-?]*[ -/]*[@-~]", "", text))


def markdown_blocks(text: str) -> list[Block]:
    from markdown_it import MarkdownIt

    tokens = MarkdownIt("commonmark").enable("table").parse(clean(text))
    blocks: list[Block] = []
    lists: list[int] = []
    prefix = ""
    heading = 0
    rows: list[list[str]] | None = None
    for token in tokens:
        if token.type in ("bullet_list_open", "ordered_list_open"):
            lists.append(
                int(token.attrGet("start") or 1) if token.type == "ordered_list_open" else 0
            )
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif token.type == "list_item_open":
            prefix = "  " * (len(lists) - 1) + (f"{lists[-1]}. " if lists[-1] else "• ")
            if lists[-1]:
                lists[-1] += 1
        elif token.type == "heading_open":
            heading = int(token.tag[1])
        elif token.type == "heading_close":
            heading = 0
        elif token.type in ("fence", "code_block"):
            blocks.append(Block("code", token.content.rstrip("\n")))
        elif token.type == "table_open":
            rows = []
        elif token.type == "tr_open" and rows is not None:
            rows.append([])
        elif token.type == "table_close":
            blocks.append(Block("table", rows=rows))
            rows = None
        elif token.type == "inline":
            # Inline markup becomes editable text; images/links never cause network access.
            text = "".join(
                f"[Image: {child.content}]"
                if child.type == "image"
                else "\n"
                if child.type in ("softbreak", "hardbreak")
                else child.content
                if child.type in ("text", "code_inline", "html_inline")
                else ""
                for child in token.children or []
            )
            if rows is not None:
                rows[-1].append(text)
            else:
                blocks.append(Block("heading" if heading else "paragraph", prefix + text, heading))
            prefix = ""
        elif token.type == "html_block":
            blocks.append(Block("code", token.content))
        elif token.type == "hr":
            blocks.append(Block("paragraph", "────────"))
    return blocks


def notebook_blocks(notebook: Notebook, ledger: Any = None, blobs: Any = None) -> list[Block]:
    manifest = notebook.read()
    title = str(manifest["metadata"].get("title") or notebook.root.name)
    blocks = [Block("title", title)]
    records: dict[str, Any] = {}
    if ledger is not None:
        # Match the latest recorded execution per cell across sessions; never execute code.
        for row in ledger.rows(
            "SELECT e.* FROM executions e JOIN sessions s USING(session_id) "
            "ORDER BY s.start,s.session_id,e.ordinal"
        ):
            if row["cell_id"]:
                records[row["cell_id"]] = row
    for index, cell in enumerate(manifest["cells"], 1):
        source = notebook.source(cell)
        if cell["type"] == "markdown":
            blocks.extend(markdown_blocks(source))
            continue
        blocks.append(
            Block("label", f"Cell {index} · {'Python' if cell['type'] == 'code' else 'Raw text'}")
        )
        blocks.append(Block("code", clean(source)))
        if cell["type"] != "code" or cell["id"] not in records:
            continue
        record = records[cell["id"]]
        blocks.append(
            Block("label", f"Recorded output · execution {record['ordinal']} · {record['status']}")
        )
        if record["source_sha256"] != source_hash(source):
            blocks.append(
                Block(
                    "warning", "These outputs were recorded from an earlier version of this cell."
                )
            )
        assert ledger is not None
        for event in fold(ledger.events(record["exec_id"])):
            kind = event["kind"]
            if kind in ("stdout", "stderr"):
                text = event.get("text")
                if text is None:
                    text = blobs.get(event["blob"]).decode("utf-8", errors="replace")
                blocks.append(Block("output", clean(text)))
            elif kind == "error":
                blocks.append(Block("warning", clean(f"{event['ename']}: {event['evalue']}")))
                trace = event.get("traceback", [])
                blocks.append(
                    Block(
                        "output",
                        clean(
                            "\n".join(cast(list[str], trace))
                            if isinstance(trace, list)
                            else str(trace)
                        ),
                    )
                )
            elif kind == "truncated":
                blocks.append(
                    Block(
                        "warning",
                        "Recorded output was truncated by the runtime; some output was not retained.",
                    )
                )
            elif kind in ("display", "result"):
                bundle = event.get("bundle", {})
                image_mime = next(
                    (mime for mime in ("image/png", "image/jpeg") if mime in bundle), None
                )
                if image_mime:
                    image = payload(bundle[image_mime], blobs)
                    blocks.append(
                        Block(
                            "image",
                            data=image if isinstance(image, bytes) else base64.b64decode(image),
                        )
                    )
                elif "text/markdown" in bundle:
                    blocks.extend(markdown_blocks(str(payload(bundle["text/markdown"], blobs))))
                elif "text/plain" in bundle:
                    blocks.append(Block("output", clean(str(payload(bundle["text/plain"], blobs)))))
                else:
                    blocks.append(
                        Block(
                            "paragraph",
                            "Rich output omitted ("
                            + ", ".join(bundle)
                            + "). Export a portable bundle to retain the original data.",
                        )
                    )
    if len(blocks) == 1:
        blocks.append(Block("paragraph", "This notebook has no cells."))
    return blocks


def pdf_bytes(blocks: list[Block]) -> bytes:
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    fonts = Path(reportlab.__file__).parent / "fonts"
    system = Path(os.environ.get("WINDIR", "/nonexistent")) / "Fonts"
    for name, candidates in {
        "KLBody": [
            system / "segoeui.ttf",
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            fonts / "Vera.ttf",
        ],
        "KLMono": [
            system / "consola.ttf",
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
            fonts / "Vera.ttf",
        ],
    }.items():
        cast(Any, pdfmetrics).registerFont(
            TTFont(name, str(next(path for path in candidates if path.is_file())))
        )
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=45,
        rightMargin=45,
        topMargin=48,
        bottomMargin=48,
        title=blocks[0].text,
        author="KernelLedger",
    )
    styles = {
        "paragraph": ParagraphStyle(
            "Body", fontName="KLBody", fontSize=10, leading=15, spaceAfter=9, wordWrap="CJK"
        ),
        "code": ParagraphStyle(
            "Code",
            fontName="KLMono",
            fontSize=8,
            leading=12,
            backColor=colors.HexColor("#F2F4F6"),
            borderPadding=5,
            spaceAfter=2,
            wordWrap="CJK",
        ),
        "label": ParagraphStyle(
            "Label",
            fontName="KLBody",
            fontSize=8,
            leading=12,
            textColor=colors.HexColor("#505A62"),
            spaceBefore=12,
            spaceAfter=5,
        ),
        "warning": ParagraphStyle(
            "Warning",
            fontName="KLBody",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#A1262F"),
            spaceAfter=7,
            wordWrap="CJK",
        ),
    }
    story: list[Any] = []
    for block in blocks:
        if block.kind == "image":
            image = Image(io.BytesIO(block.data))
            ratio = min(doc.width / image.imageWidth, (doc.height - 24) / image.imageHeight, 1)
            image.drawWidth, image.drawHeight = image.imageWidth * ratio, image.imageHeight * ratio
            story.extend([image, Spacer(1, 9)])
        elif block.kind == "table":
            rows = block.rows or [[""]]
            columns = max(len(row) for row in rows)
            table = LongTable(
                [[Paragraph(escape(cell), styles["paragraph"]) for cell in row] for row in rows],
                colWidths=[doc.width / columns] * columns,
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EBEFF4")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D6DD")),
                    ]
                )
            )
            story.extend([table, Spacer(1, 9)])
        elif block.kind in ("title", "heading"):
            level = block.level or 1
            style = ParagraphStyle(
                "Heading",
                parent=styles["paragraph"],
                fontSize=22 if block.kind == "title" else max(11, 20 - level * 2),
                leading=27 if block.kind == "title" else max(16, 26 - level * 2),
                spaceBefore=12,
                spaceAfter=10,
                keepWithNext=True,
            )
            story.append(Paragraph(escape(clean(block.text)), style))
        elif block.kind in ("code", "output"):
            style = (
                styles["code"]
                if block.kind == "code"
                else ParagraphStyle("Output", parent=styles["code"], backColor=None)
            )
            # One flowable per line lets even very long cells paginate safely.
            for line in block.text.expandtabs(4).split("\n"):
                story.append(Paragraph(escape(line).replace(" ", "&#160;") or "&#160;", style))
            story.append(Spacer(1, 7))
        else:
            story.append(
                Paragraph(
                    escape(clean(block.text)).replace("\n", "<br/>"),
                    styles.get(block.kind, styles["paragraph"]),
                )
            )

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("KLBody", 8)
        canvas.setFillColor(colors.HexColor("#505A62"))
        canvas.drawString(45, 25, "KernelLedger")
        canvas.drawRightString(A4[0] - 45, 25, str(document.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def word_bytes(blocks: list[Block]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    document.core_properties.title = blocks[0].text
    document.core_properties.author = "KernelLedger"
    normal: Any = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for block in blocks:
        if block.kind == "title":
            document.add_heading(clean(block.text), 0)
        elif block.kind == "heading":
            document.add_heading(clean(block.text), min(block.level, 9))
        elif block.kind == "image":
            from PIL import Image

            with Image.open(io.BytesIO(block.data)) as image:
                width, height = image.size
            ratio = min(6 / width, 8 / height)
            document.add_picture(
                io.BytesIO(block.data), width=Inches(width * ratio), height=Inches(height * ratio)
            )
        elif block.kind == "table":
            rows = block.rows or [[""]]
            table = document.add_table(rows=0, cols=max(len(row) for row in rows))
            table.style = "Light Shading Accent 1"
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=False):
                    cell.text = clean(value)
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(clean(block.text))
            if block.kind in ("code", "output"):
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif block.kind == "label":
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("505A62")
                run.bold = True
            elif block.kind == "warning":
                run.font.color.rgb = RGBColor.from_string("A1262F")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_document(
    notebook: Notebook,
    path: str,
    format: str,
    ledger: Any = None,
    blobs: Any = None,
    markdown: str | None = None,
) -> dict[str, str]:
    if format not in ("pdf", "docx"):
        raise SourceError("invalid_request", "Document format must be pdf or docx")
    destination = Path(path)
    if destination.suffix.lower() != f".{format}":
        raise SourceError("invalid_request", f"Choose a file ending in .{format}")
    blocks = (
        notebook_blocks(notebook, ledger, blobs)
        if markdown is None
        else [Block("title", "Assistant report"), *markdown_blocks(markdown)]
    )
    atomic(destination, pdf_bytes(blocks) if format == "pdf" else word_bytes(blocks))
    return {"path": str(destination), "format": format}
